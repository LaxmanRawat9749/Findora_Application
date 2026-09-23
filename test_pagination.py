import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

def test():
    doc = docx.Document()
    section1 = doc.sections[0]
    section1.page_width = Inches(8.27)
    section1.page_height = Inches(11.69)
    section1.left_margin = Inches(1.5)
    section1.right_margin = Inches(1.0)
    section1.top_margin = Inches(1.0)
    section1.bottom_margin = Inches(1.0)

    # Set page number type to Roman for section 1
    sectPr1 = section1._sectPr
    pgNumType1 = OxmlElement('w:pgNumType')
    pgNumType1.set(qn('w:fmt'), 'roman')
    pgNumType1.set(qn('w:start'), '1')
    sectPr1.append(pgNumType1)

    # Header / Footer for section 1
    footer1 = section1.footer
    fp1 = footer1.paragraphs[0]
    fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run1 = fp1.add_run()
    fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    f_run1._r.append(fld1)
    f_run1.font.name = 'Times New Roman'
    f_run1.font.size = Pt(10)

    p = doc.add_paragraph('This is preliminary page in Roman.')
    doc.add_section(WD_SECTION.NEW_PAGE)

    section2 = doc.sections[1]
    section2.page_width = Inches(8.27)
    section2.page_height = Inches(11.69)
    section2.left_margin = Inches(1.5)
    section2.right_margin = Inches(1.0)
    section2.top_margin = Inches(1.0)
    section2.bottom_margin = Inches(1.0)

    # Unlink footer
    section2.footer.is_linked_to_previous = False

    sectPr2 = section2._sectPr
    pgNumType2 = OxmlElement('w:pgNumType')
    pgNumType2.set(qn('w:fmt'), 'decimal')
    pgNumType2.set(qn('w:start'), '1')
    sectPr2.append(pgNumType2)

    footer2 = section2.footer
    fp2 = footer2.paragraphs[0]
    fp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run2 = fp2.add_run()
    fld2 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    f_run2._r.append(fld2)
    f_run2.font.name = 'Times New Roman'
    f_run2.font.size = Pt(10)

    p2 = doc.add_paragraph('This is chapter 1 page in Arabic.')
    doc.save('test_page_num.docx')
    print('test_page_num.docx created successfully!')

if __name__ == '__main__':
    test()
