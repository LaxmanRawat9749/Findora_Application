import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Color Palette Constants
COLOR_BG = RGBColor(11, 15, 25)          # #0B0F19 - Deep Midnight Navy
COLOR_CARD_BG = RGBColor(24, 32, 54)     # #182036 - Sleek Dark Card Surface
COLOR_CARD_BORDER = RGBColor(51, 65, 85) # #334155 - Card Outline Accent
COLOR_CARD_ACCENT = RGBColor(83, 74, 183)# #534AB7 - Findora Primary Purple
COLOR_CARD_HOVER = RGBColor(37, 48, 77)  # #25304D - Alternate Card
COLOR_TEXT_PRIMARY = RGBColor(248, 250, 252) # #F8FAFC - White Primary Text
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 - Slate Muted Text
COLOR_TEXT_PURPLE = RGBColor(165, 180, 252)  # #A5B4FC - Soft Indigo Accent
COLOR_TEXT_GREEN = RGBColor(74, 222, 128)    # #4ADE80 - Success Green Accent
COLOR_TEXT_AMBER = RGBColor(251, 191, 36)    # #FBBF24 - Warning Amber Accent
COLOR_TEXT_RED = RGBColor(248, 113, 113)     # #F87171 - Alert Red Accent
COLOR_BADGE_BG = RGBColor(45, 37, 80)        # #2D2550 - Badge Background

def init_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def create_base_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Background full fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()
    return slide

def add_header(slide, topic_badge, main_title, slide_num, total_slides=15):
    # Topic Badge Pill
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(3.4), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLOR_BADGE_BG
    badge.line.color.rgb = COLOR_CARD_ACCENT
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.word_wrap = False
    tf_b.margin_top = tf_b.margin_bottom = tf_b.margin_left = tf_b.margin_right = 0
    p_b = tf_b.paragraphs[0]
    p_b.text = topic_badge.upper()
    p_b.font.name = "Calibri"
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_TEXT_PURPLE
    p_b.alignment = PP_ALIGN.CENTER

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(9.5), Inches(0.6))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_top = tf_t.margin_bottom = tf_t.margin_left = tf_t.margin_right = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = main_title
    p_t.font.name = "Calibri"
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_PRIMARY

    # Right Header (Project Info & Slide Number)
    header_right = slide.shapes.add_textbox(Inches(9.0), Inches(0.45), Inches(3.5), Inches(0.5))
    tf_hr = header_right.text_frame
    tf_hr.margin_top = tf_hr.margin_bottom = tf_hr.margin_left = tf_hr.margin_right = 0
    p_hr = tf_hr.paragraphs[0]
    p_hr.text = f"FINDORA  •  BCA 8TH SEMESTER DEFENSE"
    p_hr.font.name = "Calibri"
    p_hr.font.size = Pt(10)
    p_hr.font.bold = True
    p_hr.font.color.rgb = COLOR_TEXT_MUTED
    p_hr.alignment = PP_ALIGN.RIGHT

    p_num = tf_hr.add_paragraph()
    p_num.text = f"{slide_num:02d} / {total_slides:02d}"
    p_num.font.name = "Calibri"
    p_num.font.size = Pt(12)
    p_num.font.bold = True
    p_num.font.color.rgb = COLOR_TEXT_PURPLE
    p_num.alignment = PP_ALIGN.RIGHT

    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Pt(1))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_CARD_BORDER
    line.line.fill.background()

def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def add_bullet_item(tf, bold_label, text_body, label_color=COLOR_TEXT_PRIMARY, body_color=COLOR_TEXT_MUTED, font_size=13, space_after=10):
    p = tf.add_paragraph() if tf.paragraphs and tf.paragraphs[0].text.strip() else tf.paragraphs[0]
    p.space_after = Pt(space_after)
    p.font.name = "Calibri"
    
    r_bullet = p.add_run()
    r_bullet.text = "• "
    r_bullet.font.bold = True
    r_bullet.font.size = Pt(font_size)
    r_bullet.font.color.rgb = COLOR_CARD_ACCENT

    r_bold = p.add_run()
    r_bold.text = bold_label + ": "
    r_bold.font.bold = True
    r_bold.font.size = Pt(font_size)
    r_bold.font.color.rgb = label_color

    r_desc = p.add_run()
    r_desc.text = text_body
    r_desc.font.size = Pt(font_size)
    r_desc.font.color.rgb = body_color

def build_presentation():
    prs = init_presentation()

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Cover)
    # =========================================================================
    s1 = create_base_slide(prs)
    
    # Left Hero Container Card
    c_hero = add_card(s1, 0.8, 0.8, 7.6, 5.9, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_ACCENT)
    tf_h = c_hero.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = tf_h.margin_right = tf_h.margin_top = Inches(0.4)
    
    p_tag = tf_h.paragraphs[0]
    p_tag.text = "POKHARA UNIVERSITY  •  BCA 8TH SEMESTER PROJECT DEFENSE"
    p_tag.font.name = "Calibri"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_TEXT_PURPLE

    p_main = tf_h.add_paragraph()
    p_main.text = "FINDORA"
    p_main.font.name = "Calibri"
    p_main.font.size = Pt(44)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_TEXT_PRIMARY
    p_main.space_before = Pt(8)

    p_sub = tf_h.add_paragraph()
    p_sub.text = "Community-Driven Lost and Found Platform"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_TEXT_GREEN
    p_sub.space_after = Pt(16)

    p_desc = tf_h.add_paragraph()
    p_desc.text = "A secure, verified digital ecosystem uniting item owners, honest finders, and platform administrators with secret question claim verification, real-time in-app chat, dual-confirmation return handshakes, and karma reputation incentives."
    p_desc.font.name = "Calibri"
    p_desc.font.size = Pt(12)
    p_desc.font.color.rgb = COLOR_TEXT_MUTED
    p_desc.space_after = Pt(20)

    # Details divider inside hero
    p_meta1 = tf_h.add_paragraph()
    p_meta1.text = "College: Oxford College of Engineering and Management, Gaidakot"
    p_meta1.font.name = "Calibri"
    p_meta1.font.size = Pt(11.5)
    p_meta1.font.color.rgb = COLOR_TEXT_PRIMARY

    p_meta2 = tf_h.add_paragraph()
    p_meta2.text = "Project Supervisors: Mr. Anil Thapaliya  |  Mr. Shiva Pathak"
    p_meta2.font.name = "Calibri"
    p_meta2.font.size = Pt(11.5)
    p_meta2.font.bold = True
    p_meta2.font.color.rgb = COLOR_TEXT_PURPLE

    p_meta3 = tf_h.add_paragraph()
    p_meta3.text = "Submitted By: Bibek Poudel (22530187)  •  Pratikshya Maske (22530255)  •  Sumina Pokhrel (22530256)"
    p_meta3.font.name = "Calibri"
    p_meta3.font.size = Pt(11)
    p_meta3.font.color.rgb = COLOR_TEXT_MUTED
    p_meta3.space_before = Pt(4)

    # Right Visual Preview Card
    c_img = add_card(s1, 8.7, 0.8, 3.8, 5.9, bg_color=COLOR_CARD_HOVER, border_color=COLOR_CARD_BORDER)
    if os.path.exists("finder dashboards 5.png"):
        s1.shapes.add_picture("finder dashboards 5.png", Inches(9.2), Inches(1.1), width=Inches(2.8))

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT
    # =========================================================================
    s2 = create_base_slide(prs)
    add_header(s2, "Topic 01 • Problem Statement", "The Crisis of Lost Valuables & Broken Traditional Channels", 2)

    # 3 Column Cards
    col_w = 3.65
    gap = 0.38
    
    # Card 1: Fragmented Channels
    c1 = add_card(s2, 0.8, 1.7, col_w, 5.2, border_color=COLOR_TEXT_RED)
    tf1 = c1.text_frame
    tf1.margin_left = tf1.margin_right = tf1.margin_top = Inches(0.25)
    p = tf1.paragraphs[0]
    p.text = "❌ Fragmented & Chaotic Channels"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_RED
    p.space_after = Pt(12)
    add_bullet_item(tf1, "Buried Social Posts", "Posts in Facebook or Viber groups are quickly buried by new feeds within minutes.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf1, "Physical Desks Siloed", "Lost items at university departments or reception desks remain isolated with no digital registry.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf1, "No Search Mechanism", "Victims cannot filter by category, date, or location, making recovery heavily dependent on luck.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)

    # Card 2: False Claims & Fraud
    c2 = add_card(s2, 0.8 + col_w + gap, 1.7, col_w, 5.2, border_color=COLOR_TEXT_AMBER)
    tf2 = c2.text_frame
    tf2.margin_left = tf2.margin_right = tf2.margin_top = Inches(0.25)
    p = tf2.paragraphs[0]
    p.text = "⚠️ False Claims & Identity Fraud"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_AMBER
    p.space_after = Pt(12)
    add_bullet_item(tf2, "Zero Verification", "Traditional methods allow anyone to falsely claim high-value items without ownership proof.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf2, "Privacy Exposure", "Owners and finders are forced to post personal phone numbers publicly, attracting spam and scams.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf2, "Unverifiable Handover", "No records exist to prove whether an item was successfully returned or stolen.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)

    # Card 3: Lack of Finder Motivation
    c3 = add_card(s2, 0.8 + (col_w + gap)*2, 1.7, col_w, 5.2, border_color=COLOR_CARD_ACCENT)
    tf3 = c3.text_frame
    tf3.margin_left = tf3.margin_right = tf3.margin_top = Inches(0.25)
    p = tf3.paragraphs[0]
    p.text = "📉 Lack of Finder Motivation"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    p.space_after = Pt(12)
    add_bullet_item(tf3, "High Effort, Zero Reward", "Finders spend personal time searching for owners with no social recognition or reward.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf3, "Social Friction", "Awkward cold calls and meeting strangers without an in-app secure coordination channel.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf3, "Passive Neglect", "Valuable items remain unreturned in drawers because finding the owner is too tedious.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)

    # =========================================================================
    # SLIDE 3: PROPOSED SOLUTION
    # =========================================================================
    s3 = create_base_slide(prs)
    add_header(s3, "Topic 02 • Proposed Solution", "Findora: The Smart, Verified Lost & Found Platform", 3)

    # Left 4 Pillars Card
    c_sol = add_card(s3, 0.8, 1.7, 7.5, 5.2)
    tf_sol = c_sol.text_frame
    tf_sol.margin_left = tf_sol.margin_right = tf_sol.margin_top = Inches(0.3)
    p = tf_sol.paragraphs[0]
    p.text = "4 Core Solution Pillars"
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    p.space_after = Pt(14)

    add_bullet_item(tf_sol, "1. Centralized Digital Repository", "Categorized feed with instant keyword search, category chips, status tags (LOST/FOUND), and GPS map pinning.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)
    add_bullet_item(tf_sol, "2. Secret Question Proof Verification", "Finders configure a secret verification question (e.g. lockscreen wallpaper, serial digit) that only the authentic owner can answer.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)
    add_bullet_item(tf_sol, "3. Dual-Confirmation Handshake", "Eliminates false recovery claims. Both Owner AND Finder must tap 'Confirm Return' in-app during physical handover.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)
    add_bullet_item(tf_sol, "4. Gamified Reputation System", "Rewards honest finders with +50 Karma Points, 5-Star Ratings, and Tiered Badges (Good Samaritan, Top Finder, Trusted Finder).", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)

    # Right Image Card
    c_img3 = add_card(s3, 8.6, 1.7, 3.9, 5.2, bg_color=COLOR_CARD_HOVER)
    if os.path.exists("finder dashboards 5.png"):
        s3.shapes.add_picture("finder dashboards 5.png", Inches(9.1), Inches(1.9), width=Inches(2.8))

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # =========================================================================
    s4 = create_base_slide(prs)
    add_header(s4, "Topic 03 • System Architecture", "Multi-Tiered Architectural Blueprint & Data Flow", 4)

    # Left Tech Flow Card
    c_arch = add_card(s4, 0.8, 1.7, 5.2, 5.2)
    tf_a = c_arch.text_frame
    tf_a.margin_left = tf_a.margin_right = tf_a.margin_top = Inches(0.3)
    p = tf_a.paragraphs[0]
    p.text = "Architectural Layers"
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    p.space_after = Pt(12)

    add_bullet_item(tf_a, "Presentation Tier", "Native Android Java client with Material 3 UI, ViewBinding, and 0ms in-memory cache.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf_a, "Application Tier", "Django REST Framework backend running on Gunicorn WSGI server handling stateless JSON APIs.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf_a, "Data Tier", "PostgreSQL relational database normalized to 3NF ensuring ACID-compliant transactional integrity.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf_a, "Cloud Storage & Media", "Backblaze B2 Object Storage (S3-compatible) storing encrypted high-resolution photos.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf_a, "Authentication & OTP", "HMAC-SHA256 JWT bearer tokens & Brevo transactional email OTP gateway.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)

    # Right Architecture Diagram
    c_diag = add_card(s4, 6.3, 1.7, 6.2, 5.2, bg_color=COLOR_CARD_HOVER)
    diag_path = "extracted_diagrams/img_9.png"
    if os.path.exists(diag_path):
        s4.shapes.add_picture(diag_path, Inches(6.5), Inches(1.9), width=Inches(5.8))

    # =========================================================================
    # SLIDE 5: CONCEPTUAL MODELS & ACTORS
    # =========================================================================
    s5 = create_base_slide(prs)
    add_header(s5, "Topic 04 • Conceptual Models", "System Actors, Roles & Use Case Flow", 5)

    # Left Actors Card
    c_act = add_card(s5, 0.8, 1.7, 5.2, 5.2)
    tf_act = c_act.text_frame
    tf_act.margin_left = tf_act.margin_right = tf_act.margin_top = Inches(0.3)
    p = tf_act.paragraphs[0]
    p.text = "3 Primary User Roles"
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    p.space_after = Pt(12)

    add_bullet_item(tf_act, "👤 Item Owner", "Posts lost item reports, searches feed, answers secret proof questions, confirms return, and rates finders.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 11)
    add_bullet_item(tf_act, "🔍 Item Finder", "Posts discovered items with mandatory photo, sets secret ownership question, chats in-app, confirms handover, and earns Karma Points.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 11)
    add_bullet_item(tf_act, "🛡️ Platform Administrator", "Moderates reported items, inspects evidence photos, resolves disputed claims, oversees user trust scores, and audits transactions.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 11)

    # Right Use Case Diagram
    c_uc = add_card(s5, 6.3, 1.7, 6.2, 5.2, bg_color=COLOR_CARD_HOVER)
    uc_path = "extracted_diagrams/img_0.png"
    if os.path.exists(uc_path):
        s5.shapes.add_picture(uc_path, Inches(6.5), Inches(1.9), width=Inches(5.8))

    # =========================================================================
    # SLIDE 6: TECHNOLOGY STACK
    # =========================================================================
    s6 = create_base_slide(prs)
    add_header(s6, "Topic 05 • Technology Stack", "Proven, Modern & Scalable Technology Stack", 6)

    cw = 5.65
    ch = 2.45
    
    # Card 1: Frontend
    c_fe = add_card(s6, 0.8, 1.7, cw, ch)
    tf = c_fe.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "📱 Native Android Frontend"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    add_bullet_item(tf, "Java 17 (Native)", "Native performance, zero bridge overhead, 60 FPS UI.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "ViewBinding & Material 3", "Type-safe view access, modern dark theme components.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Retrofit 2 & OkHttp 3", "Type-safe asynchronous REST networking with token interceptors.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Glide Image Loader", "Hardware bitmap decoding with L1 memory & L2 disk caching.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)

    # Card 2: Backend
    c_be = add_card(s6, 6.85, 1.7, cw, ch)
    tf = c_be.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "⚙️ Backend & API Services"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    add_bullet_item(tf, "Python 3.12 & Django 5.1/6.0", "Robust enterprise web framework with built-in ORM.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Django REST Framework", "Stateless JSON serializers, viewsets, and permissions.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "SimpleJWT Authentication", "Stateless HMAC-SHA256 access & refresh bearer tokens.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Gunicorn & Whitenoise", "High-concurrency WSGI application server with static asset engine.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)

    # Card 3: Database & Cloud
    c_db = add_card(s6, 0.8, 4.45, cw, ch)
    tf = c_db.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "🗄️ Database & Cloud Infrastructure"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_AMBER
    add_bullet_item(tf, "PostgreSQL Database", "ACID-compliant relational database normalized to 3NF.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Backblaze B2 Cloud Storage", "S3-compatible encrypted cloud bucket for user photos.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Brevo (Sendinblue) API", "High-deliverability transactional email service for OTP verification.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)

    # Card 4: DevOps & Tools
    c_dev = add_card(s6, 6.85, 4.45, cw, ch)
    tf = c_dev.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "🛠️ Payments & DevOps Pipeline"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    add_bullet_item(tf, "eSewa Payment Gateway", "Digital wallet integration for optional paid featured listings.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Render Cloud Platform", "Managed automated continuous deployment with health checks.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)
    add_bullet_item(tf, "Git, GitHub & Postman", "Branch-based team version control and automated API test collections.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 11.5, 4)

    # =========================================================================
    # SLIDE 7: APP WALKTHROUGH - DISCOVERY & DETAILS
    # =========================================================================
    s7 = create_base_slide(prs)
    add_header(s7, "Topic 06 • App Walkthrough", "Mobile Experience: Discovery, Search & Item Details", 7)

    # Left Highlights Card
    c_w1 = add_card(s7, 0.8, 1.7, 5.2, 5.2)
    tf = c_w1.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Key UI & Functional Features"
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    p.space_after = Pt(12)

    add_bullet_item(tf, "Instant Live Search", "Real-time search filtering listings dynamically by title, keyword, or landmark.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Category Filter Chips", "Quick toggles for Electronics, Wallets, Documents, Keys, Books, and Others.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Status Badges", "High-contrast visual tags indicating 'LOST' (red) and 'FOUND' (green) states.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Google Maps Pinning", "'View Map' button opening external Google Maps at exact item coordinates.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Contact & Claim Actions", "Direct triggers for 1-on-1 private chat and secret question verification.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)

    # Right 2 Screenshots
    c_img7 = add_card(s7, 6.3, 1.7, 6.2, 5.2, bg_color=COLOR_CARD_HOVER)
    if os.path.exists("finder dashboards 5.png"):
        s7.shapes.add_picture("finder dashboards 5.png", Inches(6.6), Inches(1.85), width=Inches(2.6))
    if os.path.exists("finderac 4.png"):
        s7.shapes.add_picture("finderac 4.png", Inches(9.5), Inches(1.85), width=Inches(2.6))

    # =========================================================================
    # SLIDE 8: APP WALKTHROUGH - REPORTING LOST & FOUND
    # =========================================================================
    s8 = create_base_slide(prs)
    add_header(s8, "Topic 07 • App Walkthrough", "Mobile Experience: Lost & Found Item Reporting", 8)

    # Left Highlights Card
    c_w2 = add_card(s8, 0.8, 1.7, 5.2, 5.2)
    tf = c_w2.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Structured Reporting Workflows"
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    p.space_after = Pt(12)

    add_bullet_item(tf, "Report Lost Item", "Capture category, distinct descriptors, date/time lost, location context, and optional photos.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Report Found Item", "Mandatory clear photo capture with camera/gallery image compression.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Secret Proof Question", "Finder sets a custom validation question (e.g. 'What is written on the back of the card?') to stop false claims.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "0ms Cache Sync", "Instant in-memory cache update ensures newly created items appear on the feed without latency.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)

    # Right 2 Screenshots
    c_img8 = add_card(s8, 6.3, 1.7, 6.2, 5.2, bg_color=COLOR_CARD_HOVER)
    if os.path.exists("report lost item.png"):
        s8.shapes.add_picture("report lost item.png", Inches(6.6), Inches(1.85), width=Inches(2.6))
    if os.path.exists("reporting found item 6.png"):
        s8.shapes.add_picture("reporting found item 6.png", Inches(9.5), Inches(1.85), width=Inches(2.6))

    # =========================================================================
    # SLIDE 9: APP WALKTHROUGH - CLAIM, CHAT & HANDSHAKE
    # =========================================================================
    s9 = create_base_slide(prs)
    add_header(s9, "Topic 08 • App Walkthrough", "Mobile Experience: Claiming, Real-Time Chat & Return", 9)

    # Left Highlights Card
    c_w3 = add_card(s9, 0.8, 1.7, 5.2, 5.2)
    tf = c_w3.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Safe Communication & Handover"
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    p.space_after = Pt(12)

    add_bullet_item(tf, "Secret Proof Claiming", "Claimant submits private answer to secret question along with optional receipts or ID proof.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "1-on-1 In-App Messaging", "Real-time chat with sent/received bubbles, photo attachment, and edit/delete context options.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Dual Confirmation Handshake", "Both Owner & Finder tap 'Confirm Return' in-app during physical handover to resolve item.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)
    add_bullet_item(tf, "Karma Points & Rating", "Finder automatically receives +50 Karma Points and public 5-Star Rating upon resolution.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 10)

    # Right 2 Screenshots
    c_img9 = add_card(s9, 6.3, 1.7, 6.2, 5.2, bg_color=COLOR_CARD_HOVER)
    if os.path.exists("report item 3.png"):
        s9.shapes.add_picture("report item 3.png", Inches(6.6), Inches(1.85), width=Inches(2.6))
    if os.path.exists("chat.png"):
        s9.shapes.add_picture("chat.png", Inches(9.5), Inches(1.85), width=Inches(2.6))

    # =========================================================================
    # SLIDE 10: WEB ADMINISTRATION & MODERATION
    # =========================================================================
    s10 = create_base_slide(prs)
    add_header(s10, "Topic 09 • Admin Console", "Web Administrator: Moderation, Dispute & Analytics", 10)

    # Left Highlights Card
    c_w4 = add_card(s10, 0.8, 1.7, 4.8, 5.2)
    tf = c_w4.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Centralized Platform Oversight"
    p.font.name = "Calibri"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_AMBER
    p.space_after = Pt(12)

    add_bullet_item(tf, "Item Moderation Registry", "Comprehensive table of all reports with filters (Pending, Approved, Flagged, Resolved).", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 8)
    add_bullet_item(tf, "Evidence & Proof Audit", "Side-by-side inspection of found item photos, claimant proofs, and secret answers.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 8)
    add_bullet_item(tf, "Dispute Resolution Override", "Administrative controls to mediate conflicting claims and prevent fraudulent takeovers.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 8)
    add_bullet_item(tf, "Platform Health & Analytics", "Real-time metrics on registered users, active reports, return rates, and server latency.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 8)

    # Right 2 Admin Screenshots stacked
    c_img10 = add_card(s10, 5.9, 1.7, 6.6, 5.2, bg_color=COLOR_CARD_HOVER)
    if os.path.exists("verify item1.png"):
        s10.shapes.add_picture("verify item1.png", Inches(6.1), Inches(1.85), width=Inches(6.2))
    if os.path.exists("Admin panel.png"):
        s10.shapes.add_picture("Admin panel.png", Inches(6.1), Inches(4.35), width=Inches(6.2))

    # =========================================================================
    # SLIDE 11: INNOVATIONS & USPs
    # =========================================================================
    s11 = create_base_slide(prs)
    add_header(s11, "Topic 10 • Innovations & USPs", "Key Architectural Innovations & Competitive Edge", 11)

    # 4 Innovation Cards Grid (2x2)
    col_w = 5.65
    row_h = 2.45

    # Card 1: Linked Parent-Child Reports
    c1 = add_card(s11, 0.8, 1.7, col_w, row_h)
    tf1 = c1.text_frame
    tf1.margin_left = tf1.margin_right = tf1.margin_top = Inches(0.25)
    p = tf1.paragraphs[0]
    p.text = "🔄 Linked Parent-Child Reports"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    add_bullet_item(tf1, "Direct Association", "Finders report directly against an owner's lost listing, creating an immediate linked match dialogue.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)
    add_bullet_item(tf1, "Zero Duplication", "Prevents multiple disconnected reports for the exact same lost property.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)

    # Card 2: Dual Confirmation Handover
    c2 = add_card(s11, 6.85, 1.7, col_w, row_h)
    tf2 = c2.text_frame
    tf2.margin_left = tf2.margin_right = tf2.margin_top = Inches(0.25)
    p = tf2.paragraphs[0]
    p.text = "🤝 Dual-Confirmation Return Handshake"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    add_bullet_item(tf2, "Mutual Verification", "Both Owner AND Finder must tap 'Confirm Return' in-app during physical handover.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)
    add_bullet_item(tf2, "Eliminates False Claims", "Guarantees no party can falsely claim an item was returned or stolen.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)

    # Card 3: Karma & Reputation System
    c3 = add_card(s11, 0.8, 4.45, col_w, row_h)
    tf3 = c3.text_frame
    tf3.margin_left = tf3.margin_right = tf3.margin_top = Inches(0.25)
    p = tf3.paragraphs[0]
    p.text = "🏆 Gamified Karma Reputation"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_AMBER
    add_bullet_item(tf3, "Immutable Ledger", "Earn +50 Karma Points for every verified return, building a transparent public score.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)
    add_bullet_item(tf3, "Tiered Badges", "Unlock Good Samaritan, Top Finder, and Trusted Finder status badges.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)

    # Card 4: Zero-Exposure Private Chat
    c4 = add_card(s11, 6.85, 4.45, col_w, row_h)
    tf4 = c4.text_frame
    tf4.margin_left = tf4.margin_right = tf4.margin_top = Inches(0.25)
    p = tf4.paragraphs[0]
    p.text = "🔒 Zero-Exposure Private In-App Chat"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    add_bullet_item(tf4, "Privacy Protection", "Coordinate safe handover meetups without sharing private phone numbers or personal emails.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)
    add_bullet_item(tf4, "Rich Controls", "Real-time message status, image sharing, edit message, and delete options.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12, 4)

    # =========================================================================
    # SLIDE 12: MONETIZATION & REVENUE MODEL
    # =========================================================================
    s12 = create_base_slide(prs)
    add_header(s12, "Topic 11 • Monetization Model", "Ethical Revenue Model: Free Utility + Paid Featured Listings", 12)

    cw = 5.65
    ch = 5.2
    
    # Left Card: Free Core
    c_free = add_card(s12, 0.8, 1.7, cw, ch, border_color=COLOR_TEXT_GREEN)
    tf = c_free.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "🆓 100% Free Core Functionality"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    p.space_after = Pt(12)

    add_bullet_item(tf, "Free Lost & Found Posting", "Anyone can report missing or discovered goods at zero cost.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)
    add_bullet_item(tf, "Free In-App Chat", "Unlimited 1-on-1 private messaging for all active items.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)
    add_bullet_item(tf, "Free Search & Maps", "Full access to category search, status filters, and GPS location.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)
    add_bullet_item(tf, "Free Verification & Handshake", "Dual-confirmation return flow is always completely free.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 12)

    # Right Card: Optional Paid Featured
    c_paid = add_card(s12, 6.85, 1.7, cw, ch, border_color=COLOR_TEXT_PURPLE)
    tf = c_paid.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "🚀 Optional Paid Featured Listings"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    p.space_after = Pt(12)

    add_bullet_item(tf, "Top Feed Placement", "Pinned to the very top of home feed and search queries for maximum visibility.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Urgent Recovery Boost", "Ideal for critical personal items (passports, laptops, wallets, keys).", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Visual 'FEATURED' Badge", "Highlighted purple badge instantly catches community finder attention.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "eSewa Wallet Integration", "Instant, secure digital payment checkout with server-side token verification.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)

    # =========================================================================
    # SLIDE 13: TESTING & QUALITY ASSURANCE
    # =========================================================================
    s13 = create_base_slide(prs)
    add_header(s13, "Topic 12 • Testing & QA", "Comprehensive Testing Strategy & Verification Results", 13)

    col_w = 3.65
    gap = 0.38

    # Card 1: Automated Tests
    c1 = add_card(s13, 0.8, 1.7, col_w, 5.2, border_color=COLOR_TEXT_GREEN)
    tf1 = c1.text_frame
    tf1.margin_left = tf1.margin_right = tf1.margin_top = Inches(0.25)
    p = tf1.paragraphs[0]
    p.text = "🧪 Automated Test Suite"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    p.space_after = Pt(12)
    add_bullet_item(tf1, "100% Pass Rate", "45 Automated Unit & Integration tests executed successfully via Django test runner.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf1, "Auth & Security Tests", "JWT token lifecycle, password hashing, and role permission enforcement.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf1, "Workflow Tests", "Item creation, claim submission, secret answer checks, and dual confirmation.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)

    # Card 2: Performance Benchmarks
    c2 = add_card(s13, 0.8 + col_w + gap, 1.7, col_w, 5.2, border_color=COLOR_TEXT_PURPLE)
    tf2 = c2.text_frame
    tf2.margin_left = tf2.margin_right = tf2.margin_top = Inches(0.25)
    p = tf2.paragraphs[0]
    p.text = "⚡ Performance & Scale"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    p.space_after = Pt(12)
    add_bullet_item(tf2, "< 180ms Latency", "Average API server response time under 180 ms for all standard REST endpoints.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf2, "200+ Concurrency", "Stress-tested with up to 200 concurrent active users with zero connection drops.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf2, "Zero DB Deadlocks", "ACID transactions maintained zero locking contention during simultaneous claims.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)

    # Card 3: Compatibility Tests
    c3 = add_card(s13, 0.8 + (col_w + gap)*2, 1.7, col_w, 5.2, border_color=COLOR_TEXT_AMBER)
    tf3 = c3.text_frame
    tf3.margin_left = tf3.margin_right = tf3.margin_top = Inches(0.25)
    p = tf3.paragraphs[0]
    p.text = "📱 Device Compatibility"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_AMBER
    p.space_after = Pt(12)
    add_bullet_item(tf3, "Physical Devices", "Smooth touch gestures, camera capture, and push alerts across Android 7.0 to 15.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf3, "Emulator & PC Mode", "Full mouse hover, 3-dots quick actions, and right-click context menu on BlueStacks 5.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)
    add_bullet_item(tf3, "Offline Cache", "0ms immediate feed rendering from L1 memory & L2 disk cache on weak networks.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 12.5, 12)

    # =========================================================================
    # SLIDE 14: CONCLUSION & FUTURE SCOPE
    # =========================================================================
    s14 = create_base_slide(prs)
    add_header(s14, "Topic 13 • Conclusion & Roadmap", "Project Summary & Future Innovation Roadmap", 14)

    cw = 5.65
    ch = 5.2
    
    # Left Card: Achievements
    c_ach = add_card(s14, 0.8, 1.7, cw, ch, border_color=COLOR_TEXT_GREEN)
    tf = c_ach.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "🎯 Key Project Achievements"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_GREEN
    p.space_after = Pt(12)

    add_bullet_item(tf, "Full-Stack System Delivered", "Fully functional native Android app, Django REST API, and PostgreSQL database.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Structured Recovery Workflow", "Replaced chaotic social media posts with categorized, geo-tagged, and searchable listings.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Fraud-Proof Security", "Engineered secret proof questions and dual-confirmation return handshakes.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Community Incentives", "Implemented gamified Karma Points (+50 pts), Tiered Badges, and 5-Star Ratings.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)

    # Right Card: Future Roadmap
    c_fut = add_card(s14, 6.85, 1.7, cw, ch, border_color=COLOR_TEXT_PURPLE)
    tf = c_fut.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "🔮 Future Innovation Roadmap"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PURPLE
    p.space_after = Pt(12)

    add_bullet_item(tf, "AI Visual Image Matching", "Automated image similarity matching using deep learning visual embeddings.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Interactive Geofencing Alerts", "Radius-based background push notifications when lost items are reported nearby.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Dynamic QR Item Tags", "Printable QR stickers for laptops/wallets enabling instant scan-to-report for finders.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)
    add_bullet_item(tf, "Multilingual Localization", "Complete Nepali language support to broaden community adoption across Nepal.", COLOR_TEXT_PRIMARY, COLOR_TEXT_MUTED, 13, 10)

    # =========================================================================
    # SLIDE 15: THANK YOU & Q&A SLIDE
    # =========================================================================
    s15 = create_base_slide(prs)
    
    # Centered Thank You Card
    c_ty = add_card(s15, 1.8, 1.2, 9.733, 5.1, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_ACCENT)
    tf_ty = c_ty.text_frame
    tf_ty.word_wrap = True
    tf_ty.margin_left = tf_ty.margin_right = tf_ty.margin_top = Inches(0.5)

    p_badge = tf_ty.paragraphs[0]
    p_badge.text = "FINDORA  •  COMMUNITY-DRIVEN LOST AND FOUND PLATFORM"
    p_badge.font.name = "Calibri"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_TEXT_PURPLE
    p_badge.alignment = PP_ALIGN.CENTER

    p_ty = tf_ty.add_paragraph()
    p_ty.text = "THANK YOU!"
    p_ty.font.name = "Calibri"
    p_ty.font.size = Pt(48)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_TEXT_PRIMARY
    p_ty.alignment = PP_ALIGN.CENTER
    p_ty.space_before = Pt(10)

    p_qa = tf_ty.add_paragraph()
    p_qa.text = "Questions, Suggestions & Feedback are Welcome"
    p_qa.font.name = "Calibri"
    p_qa.font.size = Pt(20)
    p_qa.font.bold = True
    p_qa.font.color.rgb = COLOR_TEXT_GREEN
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.space_before = Pt(8)
    p_qa.space_after = Pt(24)

    p_cred = tf_ty.add_paragraph()
    p_cred.text = "Bibek Poudel  •  Pratikshya Maske  •  Sumina Pokhrel\nSupervised By: Mr. Anil Thapaliya  |  Mr. Shiva Pathak\nOxford College of Engineering and Management, Gaidakot, Nawalpur"
    p_cred.font.name = "Calibri"
    p_cred.font.size = Pt(12)
    p_cred.font.color.rgb = COLOR_TEXT_MUTED
    p_cred.alignment = PP_ALIGN.CENTER

    # Save final presentation
    output_path = "findora_presentation.pptx"
    prs.save(output_path)
    print(f"Generated final presentation successfully: {output_path}")

if __name__ == "__main__":
    build_presentation()
